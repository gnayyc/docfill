#!/usr/bin/env python3
"""Verify the output document was processed correctly."""

from docx import Document

def verify_document(doc_path):
    """Verify document content."""
    doc = Document(doc_path)
    
    print(f"Verifying: {doc_path}")
    print("-" * 40)
    
    for i, paragraph in enumerate(doc.paragraphs):
        print(f"Paragraph {i}: '{paragraph.text}'")
        
        # Check if any placeholders remain
        if '{{' in paragraph.text and '}}' in paragraph.text:
            print(f"  ⚠️  Still contains placeholders!")
        else:
            print(f"  ✅ Processed successfully")
    
    print("-" * 40)

if __name__ == "__main__":
    print("Original template:")
    verify_document("test_template.docx")
    
    print("\nProcessed output:")
    verify_document("test_output.docx")